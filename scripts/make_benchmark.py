"""Build the multi-format benchmark corpus in data/bench/ (gitignored: third-party documents are downloaded from
their sources, not redistributed) plus bench_exact.json, questions whose answers come straight from the source data.

Formats: digital PDF (papers, a form, a long report, an invoice), scanned PDF, HTML, Markdown, CSV, XLSX, DOCX,
PPTX and PNG. Run: python scripts/make_benchmark.py   (needs network once; pip install reportlab for the invoice)
"""
import random
import re
import shutil
import sys
import urllib.request
from io import BytesIO, StringIO
from pathlib import Path

OUT = Path("data/bench")
SRC = OUT / "_sources"
UA = {"User-Agent": "VeritasBenchmark/1.0 (document QA evaluation)"}
SOURCES = {
    "adam.pdf": "https://arxiv.org/pdf/1412.6980",
    "batchnorm.pdf": "https://arxiv.org/pdf/1502.03167",
    "layernorm.pdf": "https://arxiv.org/pdf/1607.06450",
    "resnet.pdf": "https://arxiv.org/pdf/1512.03385",
    "fw4.pdf": "https://www.irs.gov/pub/irs-pdf/fw4.pdf",
    "nist_csf2.pdf": "https://nvlpubs.nist.gov/nistpubs/CSWP/NIST.CSWP.29.pdf",
    "eiffel.html": "https://en.wikipedia.org/wiki/Eiffel_Tower",
    "apollo11.html": "https://en.wikipedia.org/wiki/Apollo_11",
    "buildings.html": "https://en.wikipedia.org/wiki/List_of_tallest_buildings",
    "fomc.html": "https://www.federalreserve.gov/newsevents/pressreleases/monetary20240918a.htm",
    "mars.html": "https://nssdc.gsfc.nasa.gov/planetary/factsheet/marsfact.html",
    "fastapi_readme.md": "https://raw.githubusercontent.com/fastapi/fastapi/master/README.md",
    "sp500.csv": "https://raw.githubusercontent.com/datasets/s-and-p-500-companies/main/data/constituents.csv",
    "gdp.csv": "https://raw.githubusercontent.com/datasets/gdp/main/data/gdp.csv",
    "constitution.html": "https://www.archives.gov/founding-docs/constitution-transcript",
    "declaration.html": "https://www.archives.gov/founding-docs/declaration-transcript",
}
random.seed(11)


def fetch(name: str, cache: Path | None) -> Path:
    dst = SRC / name
    if not dst.exists():
        if cache and (cache / name).exists():
            shutil.copy(cache / name, dst)
        else:
            req = urllib.request.Request(SOURCES[name], headers=UA)
            dst.write_bytes(urllib.request.urlopen(req, timeout=90).read())
    return dst


def visible_text(html: str) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "lxml")
    for t in soup(["script", "style", "nav", "footer", "header"]):
        t.decompose()
    return re.sub(r"\s+", " ", soup.get_text(" "))


def page_image(lines: list[str], title: str):
    """Render text lines as a slightly rotated, noisy 'scanned' page image (no text layer)."""
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    w, h = 1240, 1754
    img = Image.new("L", (w, h), 248)
    d = ImageDraw.Draw(img)
    try:
        bold, reg = ImageFont.truetype("timesbd.ttf", 34), ImageFont.truetype("times.ttf", 27)
    except OSError:
        bold = reg = ImageFont.load_default()
    d.text((100, 110), title, font=bold, fill=15)
    y = 190
    for line in lines:
        d.text((100, y), line, font=reg, fill=25)
        y += 40
    for _ in range(12000):
        d.point((random.randint(0, w - 1), random.randint(0, h - 1)), fill=random.randint(110, 200))
    return img.rotate(0.7, fillcolor=240, resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.6))


def wrap(text: str, width: int = 78) -> list[str]:
    import textwrap
    return textwrap.wrap(text, width)


def build(cache: Path | None):
    import pandas as pd
    OUT.mkdir(exist_ok=True)
    SRC.mkdir(exist_ok=True)
    exact = []  # (file, kind, split, question, answer)

    # ---- real documents used as-is
    for name in ["adam.pdf", "batchnorm.pdf", "layernorm.pdf", "resnet.pdf", "fw4.pdf", "nist_csf2.pdf",
                 "eiffel.html", "fomc.html", "mars.html", "fastapi_readme.md"]:
        shutil.copy(fetch(name, cache), OUT / name)

    # ---- CSV (real): S&P 500 constituents; exact lookups
    sp = pd.read_csv(fetch("sp500.csv", cache))
    sp.to_csv(OUT / "sp500_constituents.csv", index=False)
    for _, r in sp.sample(6, random_state=3).iterrows():
        exact.append(("sp500_constituents.csv", "messy", "dev", f"In the S&P 500 constituents list, what GICS sector "
                      f"is {r['Security']} ({r['Symbol']}) in, and where is it headquartered?",
                      f"{r['GICS Sector']}; headquartered in {r['Headquarters Location']}."))
    for _, r in sp.sample(4, random_state=8).iterrows():
        exact.append(("sp500_constituents.csv", "direct", "dev", f"Which company has the ticker symbol {r['Symbol']} in "
                      "the S&P 500 constituents list?", f"{r['Security']}."))

    # ---- XLSX (real data, spreadsheet format): GDP of 20 economies, wide by year; exact lookups
    gdp = pd.read_csv(fetch("gdp.csv", cache))
    names = ["United States", "China", "Japan", "Germany", "India", "United Kingdom", "France", "Italy", "Brazil",
             "Canada", "Korea, Rep.", "Australia", "Mexico", "Spain", "Indonesia", "Netherlands", "Saudi Arabia",
             "Turkiye", "Switzerland", "Poland"]
    sub = gdp[gdp["Country Name"].isin(names) & gdp["Year"].between(2015, 2022)]
    wide = sub.pivot(index="Country Name", columns="Year", values="Value").reset_index()
    wide.columns = ["Country"] + [str(c) for c in wide.columns[1:]]
    with pd.ExcelWriter(OUT / "world_gdp_2015_2022.xlsx") as xw:
        wide.round(0).to_excel(xw, sheet_name="GDP current USD", index=False)
    for _, r in wide.sample(10, random_state=5).iterrows():
        year = random.choice([str(y) for y in range(2015, 2023)])
        exact.append(("world_gdp_2015_2022.xlsx", "messy", "heldout", f"In the GDP spreadsheet, what was "
                      f"{r['Country']}'s GDP in {year} (current US dollars)?", f"{round(r[year]):,} US dollars."))

    # ---- DOCX (real text, Word format): US Constitution, Articles I-VII
    from docx import Document
    text = visible_text(fetch("constitution.html", cache).read_text(encoding="utf-8"))
    body = text[text.index("We the People"):text.index("Article. VII.") + 600]
    doc = Document()
    doc.add_heading("The Constitution of the United States (transcript)", 0)
    for part in re.split(r"(?=Article\. [IVX]+\.)|(?=Section\. \d+\.)", body):
        part = part.strip()
        if not part:
            continue
        head = re.match(r"(Article\. [IVX]+\.|Section\. \d+\.)", part)
        if head:
            doc.add_heading(head[1], 1 if head[1].startswith("Article") else 2)
            part = part[len(head[1]):].strip()
        doc.add_paragraph(part)
    doc.save(OUT / "us_constitution.docx")
    for q, a in [("According to the Constitution, how old must a Representative be?", "Twenty five years."),
                 ("How old must a Senator be, and for how many years must they have been a citizen?",
                  "Thirty years old and nine years a citizen of the United States."),
                 ("How many Senators does each State have, and how long is their term?",
                  "Two Senators from each State, for six years."),
                 ("What is the minimum age to be eligible for President?", "Thirty five years."),
                 ("How many States had to ratify the Constitution for it to be established?",
                  "The ratification of the Conventions of nine States."),
                 ("What does the Preamble say the Constitution is ordained to do first?",
                  "To form a more perfect Union.")]:
        exact.append(("us_constitution.docx", "direct", "dev", q, a))

    # ---- PPTX (real facts, slide format): Apollo 11, every fact checked against the Wikipedia text
    from pptx import Presentation
    from pptx.util import Inches, Pt
    apollo = visible_text(fetch("apollo11.html", cache).read_text(encoding="utf-8"))
    slides = [("Apollo 11 at a glance", ["First crewed Moon landing", "Launched July 16, 1969 from Kennedy Space Center",
                                         "Saturn V rocket"]),
              ("Crew", ["Commander: Neil Armstrong", "Lunar Module Pilot: Buzz Aldrin",
                        "Command Module Pilot: Michael Collins"]),
              ("Spacecraft", ["Command module: Columbia", "Lunar module: Eagle"]),
              ("Landing", ["Eagle landed July 20, 1969", "Landing site: Sea of Tranquility",
                           "Armstrong walked on the Moon about six hours after landing; Aldrin followed nineteen minutes later"]),
              ("Return", ["Splashdown in the Pacific Ocean on July 24, 1969", "Recovery ship: USS Hornet"])]
    facts = ["July 16, 1969", "Kennedy Space Center", "Saturn V", "Neil Armstrong", "Buzz Aldrin", "Michael Collins",
             "Columbia", "Eagle", "July 20, 1969", "Sea of Tranquility", "July 24, 1969",
             "Pacific Ocean", "USS Hornet", "nineteen minutes later", "six hours after landing"]
    missing = [f for f in facts if f not in apollo]  # keep the deck honest: every fact is in the source article
    assert not missing, f"not in the Apollo 11 article: {missing}"
    prs = Presentation()
    for title, bullets in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
    prs.save(OUT / "apollo11_mission.pptx")
    for q, a in [("Who were the three Apollo 11 astronauts and their roles?",
                  "Neil Armstrong (commander), Buzz Aldrin (lunar module pilot), Michael Collins (command module pilot)."),
                 ("When and from where did Apollo 11 launch?", "July 16, 1969, from Kennedy Space Center."),
                 ("What were the names of the command module and lunar module?", "Columbia and Eagle."),
                 ("Where did the lunar module land and on what date?", "The Sea of Tranquility, July 20, 1969."),
                 ("How did the mission end?", "Splashdown in the Pacific Ocean on July 24, 1969; recovered by USS Hornet."),
                 ("What rocket launched Apollo 11?", "Saturn V.")]:
        exact.append(("apollo11_mission.pptx", "direct", "heldout", q, a))

    # ---- PNG (real data, image format): 10 tallest buildings, rendered as a table image
    from PIL import Image, ImageDraw, ImageFont
    t = pd.read_html(StringIO(fetch("buildings.html", cache).read_text(encoding="utf-8")))[1]
    t.columns = [c[0] if c[0] == c[1] else f"{c[0]} {c[1]}" for c in t.columns]
    hcol = next(c for c in t.columns if c.startswith("Height") and c.endswith("m"))
    top = t[["Name", hcol, "Floors", "City", "Country", "Year"]].head(10).copy()
    top["Name"] = top["Name"].astype(str).str.replace(r"\[.*?\]", "", regex=True)
    top["Floors"] = top["Floors"].astype(str).str.split(" (", regex=False).str[0]  # "163 (+ 2 below ground)" -> "163"
    top["Year"] = top["Year"].astype(int)
    img = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(img)
    try:
        f, fb = ImageFont.truetype("arial.ttf", 24), ImageFont.truetype("arialbd.ttf", 26)
    except OSError:
        f = fb = ImageFont.load_default()
    d.text((30, 20), "World's 10 tallest buildings (source: Wikipedia)", font=fb, fill="black")
    xs = [30, 560, 720, 830, 1080, 1340]
    for x, head in zip(xs, ["Building", "Height (m)", "Floors", "City", "Country", "Year"]):
        d.text((x, 80), head, font=fb, fill="black")
    for i, (_, r) in enumerate(top.iterrows()):
        y = 130 + i * 60
        d.line((25, y - 10, 1475, y - 10), fill="#bbbbbb")
        for x, v in zip(xs, [r["Name"], r[hcol], r["Floors"], r["City"], r["Country"], r["Year"]]):
            d.text((x, y), str(v), font=f, fill="black")
    img.save(OUT / "tallest_buildings.png")
    for _, r in top.sample(6, random_state=2).iterrows():
        exact.append(("tallest_buildings.png", "messy", "heldout",
                      f"In the tallest buildings image, how tall is {r['Name']} and in which city is it?",
                      f"{r[hcol]} m; {r['City']}."))
    for _, r in top.sample(4, random_state=6).iterrows():
        exact.append(("tallest_buildings.png", "messy", "heldout",
                      f"In the tallest buildings image, how many floors does {r['Name']} have and when was it completed?",
                      f"{r['Floors']} floors; completed in {r['Year']}."))

    # ---- scanned PDF (real text, image-only pages): Declaration of Independence
    dec = visible_text(fetch("declaration.html", cache).read_text(encoding="utf-8"))
    dec = dec[dec.index("In Congress, July 4, 1776"):]
    dec = dec[:dec.index("He has refused for a long time")]
    lines = wrap(dec, 80)
    pages = [page_image(lines[i:i + 36], "The Declaration of Independence (scanned transcript)")
             for i in range(0, min(len(lines), 72), 36)]
    pages[0].save(OUT / "declaration_scan.pdf", save_all=True, append_images=pages[1:], resolution=150)
    for q, a in [("On what date was the Declaration adopted in Congress, according to the scan?", "July 4, 1776."),
                 ("How many united States of America does the Declaration's title refer to?", "Thirteen."),
                 ("According to the Declaration, what unalienable Rights are named?",
                  "Life, Liberty and the pursuit of Happiness."),
                 ("From what do governments derive their just powers, per the Declaration?",
                  "From the consent of the governed."),
                 ("Who is described as having a history of repeated injuries and usurpations?",
                  "The present King of Great Britain.")]:
        exact.append(("declaration_scan.pdf", "messy", "dev", q, a))

    # ---- invoice PDF (fictional business document, digital)
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
    from reportlab.lib.styles import getSampleStyleSheet
    items = [("Steel shelving unit, 5-tier", 12, 89.50), ("Pallet jack, 2.5 t", 2, 412.00),
             ("Safety gloves (box of 50)", 30, 14.75), ("LED high-bay light 150 W", 16, 64.20),
             ("Label printer", 1, 239.99), ("Barcode scanner, wireless", 4, 118.40)]
    sub_total = round(sum(q * p for _, q, p in items), 2)
    tax = round(sub_total * 0.0825, 2)
    ss = getSampleStyleSheet()
    rows = [["Item", "Qty", "Unit price", "Line total"]] + [[n, str(q), f"${p:,.2f}", f"${q * p:,.2f}"] for n, q, p in items]
    rows += [["", "", "Subtotal", f"${sub_total:,.2f}"], ["", "", "Sales tax 8.25%", f"${tax:,.2f}"],
             ["", "", "Total due", f"${sub_total + tax:,.2f}"]]
    SimpleDocTemplate(str(OUT / "invoice_INV-2025-0417.pdf"), pagesize=A4, leftMargin=18 * mm).build([
        Paragraph("Northbridge Supplies Ltd. (fictional)", ss["Title"]),
        Paragraph("INVOICE INV-2025-0417 - Issued 14 March 2025 - Payment due 13 April 2025 (Net 30)", ss["Normal"]),
        Paragraph("Bill to: Harbor Logistics Co., 88 Quay Street, Portland, OR", ss["Normal"]), Spacer(1, 8 * mm),
        Table(rows), Spacer(1, 6 * mm),
        Paragraph("Late payments incur 1.5% interest per month. Bank: First Coastal Bank, account 4471-0098.", ss["Normal"])])
    for q, a in [("What is the invoice number and when is payment due?", "INV-2025-0417; due 13 April 2025 (Net 30)."),
                 ("What is the total due on the invoice?", f"${sub_total + tax:,.2f}."),
                 ("What is the subtotal before tax and the tax amount?", f"${sub_total:,.2f} subtotal; ${tax:,.2f} tax (8.25%)."),
                 ("How many pallet jacks were ordered and at what unit price?", "2 at $412.00 each."),
                 ("What is the line total for the LED high-bay lights?", f"${16 * 64.20:,.2f} (16 x $64.20)."),
                 ("Who is the invoice billed to?", "Harbor Logistics Co., 88 Quay Street, Portland, OR."),
                 ("What is the late payment penalty?", "1.5% interest per month.")]:
        exact.append(("invoice_INV-2025-0417.pdf", "messy", "dev", q, a))

    import json
    Path("bench_exact.json").write_text(json.dumps(
        [dict(file=f, type=k, split=s, question=q, ground_truth=a) for f, k, s, q, a in exact], indent=1,
        ensure_ascii=False), encoding="utf-8")
    print(f"{len(list(OUT.glob('*.*')))} documents in {OUT}; {len(exact)} exact questions -> bench_exact.json")


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else None)
